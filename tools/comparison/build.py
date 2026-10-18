# -*- coding: utf-8 -*-
"""Build docs/comparison-chrome-vs-swiftra.{md,xlsx,html} from rows.py.

Usage: python3 tools/comparison/build.py
"""
import collections, html, os, sys
sys.path.insert(0, os.path.dirname(__file__))
import rows as D
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
WIN = ["Swiftra", "Chrome", "बराबर", "अनिर्णीत", "लागू नहीं"]
FILL = {"Swiftra": "C6EFCE", "Chrome": "FFC7CE", "बराबर": "DDEBF7",
        "अनिर्णीत": "FFEB9C", "लागू नहीं": "E7E6E6"}
FULL = D.full_rows()


def stats():
    total = collections.Counter(r[7] for r in FULL)
    assert not set(total) - set(WIN), total
    cats = collections.OrderedDict()
    for r in FULL:
        cats.setdefault(r[1], collections.Counter())[r[7]] += 1
    wc = ws = 0
    for r in D.R:
        c, s = D.points(r)
        wc += c; ws += s
    high = collections.Counter(r[7] for r in FULL if r[9] == "उच्च")
    return total, cats, wc, ws, high


def verdict(total, wc, ws, high):
    n = len(FULL)
    return [
        ("ईमानदार फ़ैसला",
         f"आज कुल मिलाकर असली Chrome आगे है: {n} पंक्तियों में Chrome {total['Chrome']}, Swiftra {total['Swiftra']} जीतता है; "
         f"महत्व-भार वाले अंक Chrome {wc} : Swiftra {ws} (बराबर में दोनों को अंक)। "
         f"Swiftra अभी 'हर चीज़ में Chrome से आगे' नहीं है।"),
        ("उच्च-महत्व पंक्तियाँ",
         f"उच्च महत्व की {sum(high.values())} पंक्तियों में: Chrome {high['Chrome']}, Swiftra {high['Swiftra']}, "
         f"बराबर {high['बराबर']}, अनिर्णीत {high['अनिर्णीत']}, लागू नहीं {high['लागू नहीं']}।"),
        ("Swiftra की जीत",
         "privacy और पारदर्शिता: Standard Safe Browsing में शून्य डेटा बाहर, telemetry नहीं, server experiments नहीं, "
         "startup पर कम requests, autofill form data Google को नहीं, doodle ping नहीं, हर बदलाव खुले patch में।"),
        ("Chrome की जीत",
         "security patch level (.84 बनाम .80), अपने-आप update, Safe Browsing coverage + Enhanced AI, HTTPS-by-default, "
         "spam-notification चेतावनी, Google सेवाएँ (Sync, Password Manager, Translate, Gemini, Discover, पढ़कर सुनाना, "
         "WebAPK, price tracking, push), और emulator पर start/memory (अलग version, असली फ़ोन नहीं)।"),
        ("बराबर",
         "एक ही Chromium 155 engine वाली चीज़ें: site isolation, DoH, PQ crypto, JPEG XL, tabs/layout, दूसरे password manager, "
         "cookies default, fingerprinting सुरक्षा का न होना, extensions का न होना, search default।"),
        ("अनिर्णीत",
         "असली फ़ोन पर speed, Speedometer, battery, stability, AI Mode, Incognito lock, ad filter, screen-share सुरक्षा, असली Arena — "
         "बिना माप के कोई दावा नहीं।"),
        ("सबसे तेज़ रास्ता",
         "क्रम 1-6 (कम/मध्यम मेहनत) से Chrome की 2 बढ़त (patch level, HTTPS) बराबर होंगी और 4 'बराबर/अनिर्णीत' पंक्तियाँ "
         "(cookies, fingerprinting, ad-block, search चुनाव) Swiftra की जीत बन सकती हैं — हर एक E2E से साबित होने के बाद ही।"),
    ]


def md():
    total, cats, wc, ws, high = stats()
    o = [f"# असली Chrome बनाम Swiftra: विस्तृत तुलना v3 ({D.DATE})\n",
         f"- {D.CHROME_REF}", f"- {D.SWIFTRA_BUILD}",
         "- स्थिति: **PASS (emulator)** = x64 Android emulator, **PASS (Redroid)** = arm64 Redroid, **static** = केवल स्रोत/कोड, "
         "**NOT RUN** = जाँचा नहीं, **कमी** = मौजूद नहीं।",
         "- विजेता: **Swiftra** / **Chrome** / **बराबर** / **अनिर्णीत** (माप/test नहीं) / **लागू नहीं** (कानूनी विषय)।",
         "- महत्व: उच्च = 3, मध्यम = 2, कम = 1 अंक। जीतने वाले को उतने अंक; बराबर में दोनों को; अनिर्णीत/लागू नहीं में किसी को नहीं।",
         "- यही data: `docs/comparison-chrome-vs-swiftra.xlsx` और `docs/comparison-chrome-vs-swiftra.html` (स्रोत: `tools/comparison/rows.py`)।\n",
         "## तुलना तालिका\n",
         "| " + " | ".join(D.COLUMNS) + " |", "|" + "---|" * len(D.COLUMNS)]
    for r in FULL:
        cells = [str(c).replace("|", "/") for c in r]
        cells[7] = f"**{cells[7]}**"
        o.append("| " + " | ".join(cells) + " |")
    o += ["\n## अंतिम सारांश\n",
          f"कुल **{len(FULL)}** पंक्तियाँ — Swiftra **{total['Swiftra']}**, Chrome **{total['Chrome']}**, बराबर **{total['बराबर']}**, "
          f"अनिर्णीत **{total['अनिर्णीत']}**, लागू नहीं **{total['लागू नहीं']}**। महत्व-भार वाले अंक: **Chrome {wc} : Swiftra {ws}**।\n",
          "| क्षेत्र | Swiftra | Chrome | बराबर | अनिर्णीत | लागू नहीं |", "|---|---|---|---|---|---|"]
    for c, k in cats.items():
        o.append(f"| {c} | " + " | ".join(str(k[w]) for w in WIN) + " |")
    o.append("| **कुल** | " + " | ".join(f"**{total[w]}**" for w in WIN) + " |\n")
    for a, b in verdict(total, wc, ws, high):
        o.append(f"- **{a}:** {b}")
    o += ["\n### Chrome को हराने का क्रम\n", "| क्रम | काम | असर | संबंधित पंक्ति | मेहनत |", "|---|---|---|---|---|"]
    o += ["| " + " | ".join(p) + " |" for p in D.PLAN]
    o.append("\n## स्रोत\n")
    o += [f"- **{a}** — {b}: {u}" for a, b, u in D.SOURCES]
    p = os.path.join(ROOT, "docs", "comparison-chrome-vs-swiftra.md")
    open(p, "w", encoding="utf-8").write("\n".join(o) + "\n")
    return p


def xlsx():
    total, cats, wc, ws_, high = stats()
    wb = Workbook()
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hf, hfont = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws = wb.active
    ws.title = "तुलना"
    ws["A1"] = f"असली Chrome 155 (Android) बनाम Swiftra — विस्तृत तुलना v3 ({D.DATE})"
    ws["A1"].font = Font(bold=True, size=15, color="1F3864")
    ws["A2"], ws["A3"] = D.CHROME_REF, D.SWIFTRA_BUILD
    for c in ("A2", "A3"):
        ws[c].font = Font(italic=True, size=9, color="595959")
    hr = 5
    widths = [5, 14, 24, 34, 38, 30, 14, 12, 34, 9, 20, 38, 9, 11]
    for j, h in enumerate(D.COLUMNS, 1):
        c = ws.cell(row=hr, column=j, value=h)
        c.fill, c.font, c.border, c.alignment = hf, hfont, border, center
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[hr].height = 34
    for i, r in enumerate(FULL, 1):
        rr = hr + i
        for j, v in enumerate(r, 1):
            c = ws.cell(row=rr, column=j, value=v)
            c.alignment, c.border = wrap, border
        for j in (8, 10, 13, 14):
            ws.cell(row=rr, column=j).alignment = center
        w = ws.cell(row=rr, column=8)
        w.fill, w.font = PatternFill("solid", fgColor=FILL[r[7]]), Font(bold=True)
        if r[9] == "उच्च":
            ws.cell(row=rr, column=10).font = Font(bold=True, color="C00000")
        longest = max(len(str(v)) / max(widths[k] - 2, 1) for k, v in enumerate(r))
        ws.row_dimensions[rr].height = max(30, min(160, 15 * (int(longest) + 1)))
    last = hr + len(FULL)
    t = Table(displayName="Comparison", ref=f"A{hr}:{get_column_letter(len(D.COLUMNS))}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    ws.freeze_panes = ws.cell(row=hr + 1, column=4)

    s = wb.create_sheet("सारांश")
    s["A1"] = "अंतिम सारांश"
    s["A1"].font = Font(bold=True, size=15, color="1F3864")
    for j, h in enumerate(["विजेता", "पंक्तियाँ", "उच्च-महत्व"], 1):
        c = s.cell(row=3, column=j, value=h)
        c.fill, c.font, c.border = hf, hfont, border
    for k, w in enumerate(WIN, 4):
        a = s.cell(row=k, column=1, value=w)
        a.fill, a.font = PatternFill("solid", fgColor=FILL[w]), Font(bold=True)
        s.cell(row=k, column=2, value=total[w])
        s.cell(row=k, column=3, value=high[w])
        for col in (1, 2, 3):
            s.cell(row=k, column=col).border = border
    s.cell(row=9, column=1, value="कुल").font = Font(bold=True)
    s.cell(row=9, column=2, value=len(FULL)).font = Font(bold=True)
    s.cell(row=9, column=3, value=sum(high.values())).font = Font(bold=True)
    s.cell(row=10, column=1, value="महत्व-भार अंक").font = Font(bold=True)
    s.cell(row=10, column=2, value=f"Chrome {wc} : Swiftra {ws_}").font = Font(bold=True, color="C00000")
    ch = BarChart()
    ch.type, ch.title, ch.legend = "bar", "विजेता के अनुसार पंक्तियाँ", None
    ch.add_data(Reference(s, min_col=2, min_row=3, max_row=8), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=4, max_row=8))
    ch.height, ch.width = 7, 14
    s.add_chart(ch, "E3")
    r0 = 13
    for j, h in enumerate(["क्षेत्र"] + WIN, 1):
        c = s.cell(row=r0, column=j, value=h)
        c.fill, c.font, c.border, c.alignment = hf, hfont, border, center
    for i, (cat, k) in enumerate(cats.items(), 1):
        s.cell(row=r0 + i, column=1, value=cat).border = border
        for j, w in enumerate(WIN, 2):
            c = s.cell(row=r0 + i, column=j, value=k[w])
            c.border, c.alignment = border, Alignment(horizontal="center")
            if k[w]:
                c.fill = PatternFill("solid", fgColor=FILL[w])
    rt = r0 + len(cats) + 1
    s.cell(row=rt, column=1, value="कुल").font = Font(bold=True)
    for j, w in enumerate(WIN, 2):
        c = s.cell(row=rt, column=j, value=total[w])
        c.font, c.alignment = Font(bold=True), Alignment(horizontal="center")
    rv = rt + 2
    for i, (a, b) in enumerate(verdict(total, wc, ws_, high)):
        s.cell(row=rv + i, column=1, value=a).font = Font(bold=True)
        s.merge_cells(start_row=rv + i, start_column=2, end_row=rv + i, end_column=9)
        s.cell(row=rv + i, column=2, value=b).alignment = wrap
        s.row_dimensions[rv + i].height = 50
    rp = rv + 8
    s.cell(row=rp, column=1, value="Chrome को हराने का क्रम").font = Font(bold=True, size=12, color="1F3864")
    for j, h in enumerate(["क्रम", "काम", "असर", "संबंधित पंक्ति", "मेहनत"], 1):
        c = s.cell(row=rp + 1, column=j, value=h)
        c.fill, c.font, c.border = hf, hfont, border
    for i, p in enumerate(D.PLAN, 2):
        for j, v in enumerate(p, 1):
            c = s.cell(row=rp + i, column=j, value=v)
            c.alignment, c.border = wrap, border
        s.row_dimensions[rp + i].height = 32
    for col, wdt in zip("ABCDEFGHI", [22, 44, 34, 20, 12, 12, 12, 12, 12]):
        s.column_dimensions[col].width = wdt

    so = wb.create_sheet("स्रोत")
    for j, h in enumerate(["ID", "स्रोत", "URL"], 1):
        c = so.cell(row=1, column=j, value=h)
        c.fill, c.font = hf, hfont
    for i, (a, b, u) in enumerate(D.SOURCES, 2):
        so.cell(row=i, column=1, value=a)
        so.cell(row=i, column=2, value=b).alignment = wrap
        c = so.cell(row=i, column=3, value=u)
        c.hyperlink, c.font = u, Font(color="0563C1", underline="single")
    so.column_dimensions["A"].width, so.column_dimensions["B"].width, so.column_dimensions["C"].width = 6, 60, 100
    p = os.path.join(ROOT, "docs", "comparison-chrome-vs-swiftra.xlsx")
    wb.save(p)
    return p


def page():
    total, cats, wc, ws, high = stats()
    e = html.escape
    color = {k: "#" + v for k, v in FILL.items()}
    n = len(FULL)
    bars = "".join(
        f'<div class="bar"><span class="bl">{e(w)}</span><span class="bt" style="width:{total[w]*100/n:.1f}%;background:{color[w]}"></span>'
        f'<b>{total[w]}</b></div>' for w in WIN)
    pts = (f'<div class="pts"><div><small>महत्व-भार अंक</small><br><span class="c">Chrome {wc}</span> : '
           f'<span class="s">Swiftra {ws}</span></div><div><small>उच्च-महत्व पंक्तियाँ ({sum(high.values())})</small><br>'
           + " · ".join(f"{e(w)} {high[w]}" for w in WIN) + "</div></div>")
    btns = '<button class="f on" data-w="">सभी</button>' + "".join(
        f'<button class="f" data-w="{e(w)}" style="background:{color[w]}">{e(w)} ({total[w]})</button>' for w in WIN)
    head = "".join(f"<th>{e(h)}</th>" for h in D.COLUMNS)
    body = []
    for r in FULL:
        tds = []
        for j, v in enumerate(r):
            v = e(str(v))
            if j == 7:
                tds.append(f'<td class="win" style="background:{color[r[7]]}">{v}</td>')
            elif j == 9:
                tds.append(f'<td class="imp {"hi" if r[9]=="उच्च" else ""}">{v}</td>')
            else:
                tds.append(f"<td>{v}</td>")
        body.append(f'<tr data-w="{e(r[7])}">' + "".join(tds) + "</tr>")
    cat_rows = "".join(f"<tr><td>{e(c)}</td>" + "".join(
        f'<td style="background:{color[w] if k[w] else "#fff"}">{k[w]}</td>' for w in WIN) + "</tr>" for c, k in cats.items())
    cat_rows += "<tr><th>कुल</th>" + "".join(f"<th>{total[w]}</th>" for w in WIN) + "</tr>"
    verdict_html = "".join(f"<li><b>{e(a)}:</b> {e(b)}</li>" for a, b in verdict(total, wc, ws, high))
    plan = "".join("<tr>" + "".join(f"<td>{e(x)}</td>" for x in p) + "</tr>" for p in D.PLAN)
    src = "".join(f'<li><b>{e(a)}</b> {e(b)} — <a href="{e(u)}">{e(u)}</a></li>' for a, b, u in D.SOURCES)
    doc = f"""<!doctype html><html lang="hi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Chrome बनाम Swiftra — तुलना v3</title><style>
body{{font-family:'Noto Sans Devanagari','Nirmala UI',system-ui,sans-serif;margin:0;background:#f5f7fb;color:#1b1f29}}
header{{background:#1F3864;color:#fff;padding:18px 22px}}header h1{{margin:0;font-size:22px}}header p{{margin:4px 0 0;font-size:12px;opacity:.85}}
main{{padding:16px 20px}}.cards{{display:flex;gap:16px;flex-wrap:wrap}}.card{{background:#fff;border-radius:10px;padding:14px 16px;box-shadow:0 1px 4px #0002;flex:1;min-width:300px}}
.bar{{display:flex;align-items:center;gap:8px;margin:6px 0}}.bl{{width:80px;font-weight:600}}.bt{{height:18px;border-radius:4px;display:inline-block;border:1px solid #0002}}
.pts{{display:flex;gap:24px;margin-top:10px;flex-wrap:wrap}}.pts .c{{color:#b00020;font-weight:700;font-size:20px}}.pts .s{{color:#0a7a2f;font-weight:700;font-size:20px}}
.f{{border:1px solid #0003;border-radius:16px;padding:5px 12px;margin:3px;cursor:pointer;background:#fff;font:inherit}}.f.on{{outline:3px solid #1F3864}}
.wrap{{overflow:auto;max-height:78vh;background:#fff;border-radius:10px;box-shadow:0 1px 4px #0002;margin-top:10px}}
table{{border-collapse:collapse;font-size:12.5px}}th,td{{border:1px solid #d9dde5;padding:6px 7px;vertical-align:top}}
.main th{{background:#1F3864;color:#fff;position:sticky;top:0;z-index:1}}.main td:nth-child(3){{font-weight:600;min-width:150px}}.main td:nth-child(4),.main td:nth-child(5),.main td:nth-child(12){{min-width:220px}}.main td:nth-child(9){{min-width:190px}}.main td:nth-child(11){{min-width:130px}}
td.win{{font-weight:700;text-align:center;white-space:nowrap}}td.imp{{text-align:center}}td.imp.hi{{color:#c00000;font-weight:700}}
.sum th{{background:#1F3864;color:#fff}}.sum td{{text-align:center}}.sum td:first-child{{text-align:left}}li{{margin:5px 0}}a{{color:#0563C1;word-break:break-all}}
</style></head><body><header><h1>असली Chrome 155 (Android) बनाम Swiftra — विस्तृत तुलना v3 ({e(D.DATE)})</h1>
<p>{e(D.CHROME_REF)}</p><p>{e(D.SWIFTRA_BUILD)}</p></header><main>
<div class="cards"><div class="card"><b>विजेता के अनुसार ({n} पंक्तियाँ)</b>{bars}{pts}</div>
<div class="card"><b>क्षेत्र-वार</b><table class="sum" style="width:100%;margin-top:8px"><tr><th>क्षेत्र</th>{''.join(f'<th>{e(w)}</th>' for w in WIN)}</tr>{cat_rows}</table></div></div>
<h2>तुलना तालिका</h2><div>{btns}</div>
<div class="wrap"><table class="main"><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table></div>
<h2>अंतिम सारांश</h2><div class="card"><ul>{verdict_html}</ul></div>
<h2>Chrome को हराने का क्रम</h2><div class="card"><table class="sum" style="width:100%"><tr><th>क्रम</th><th>काम</th><th>असर</th><th>संबंधित पंक्ति</th><th>मेहनत</th></tr>{plan}</table></div>
<h2>स्रोत</h2><div class="card"><ul>{src}</ul></div>
<p style="font-size:11px;color:#666">महत्व: उच्च = 3, मध्यम = 2, कम = 1 अंक; जीतने वाले को उतने अंक, बराबर में दोनों को, अनिर्णीत/लागू नहीं में किसी को नहीं।</p>
</main><script>
document.querySelectorAll('.f').forEach(b=>b.onclick=()=>{{document.querySelectorAll('.f').forEach(x=>x.classList.remove('on'));b.classList.add('on');
const w=b.dataset.w;document.querySelectorAll('.main tbody tr').forEach(tr=>tr.style.display=(!w||tr.dataset.w===w)?'':'none');}});
</script></body></html>"""
    p = os.path.join(ROOT, "docs", "comparison-chrome-vs-swiftra.html")
    open(p, "w", encoding="utf-8").write(doc)
    return p


if __name__ == "__main__":
    for f in (md, xlsx, page):
        print(f())
    total, cats, wc, ws, high = stats()
    print(len(FULL), dict(total), "weighted", wc, ws, "high", dict(high))
